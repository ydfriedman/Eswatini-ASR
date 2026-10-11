#!/usr/bin/env python3
"""Extract Palladium-staff LOE (days) by person / budget / milestone.

Reads the 'Milestone Budget' tab of every .xlsx in INPUT_DIR (section A,
"PALLADIUM STAFF") and writes an Excel workbook with:
  - LOE Table          person | budget | milestone | LOE | role | person total
  - Totals by Person   one row per person, summed across all budgets
  - Reconciliation     table days vs. each workbook's own end-of-year total column
Integrity warnings are printed and written to a 'Warnings' sheet.

Usage: python3 scripts/extract_loe.py INPUT_DIR OUTPUT.xlsx
"""
import difflib
import glob
import os
import sys
import warnings
from collections import defaultdict

import openpyxl
from openpyxl.utils import get_column_letter as col_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

warnings.filterwarnings("ignore")

SHEET = "Milestone Budget"
HEADER_ROW, CODE_ROW, FIRST_DATA_ROW = 9, 7, 11   # row 9 holds "# Days"; row 7 milestone codes
NAME_COL, ROLE_COL, BUDGET_CELL = 3, 1, "C1"       # person in C, position label in A, budget in C1
SECTION_B_MARKER = "B"                             # column B value that starts the next section
SKIP_FILES = ("HQ LOE Distribution Summary",)      # workbooks without a Milestone Budget tab
EXCLUDE_NAME_PREFIXES = ("TBD",)                   # placeholder staff, dropped from the table
SUBTOTAL_PREFIX = "SUBTOTAL"                       # subtotal row, never a person
# Rows with a blank name cell whose position label contains one of these
# strings are attributed to that person instead of being dropped.
UNNAMED_ROW_OVERRIDES = {"Ogana, Nelson": "Ogana, Nelson"}


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def extract(path, warns):
    wbv = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wbv.sheetnames:
        warns.append((os.path.basename(path), "no 'Milestone Budget' tab - skipped"))
        return None
    ws = wbv[SHEET]
    wf = openpyxl.load_workbook(path)[SHEET]  # formulas, to detect uncached values
    fname = os.path.basename(path)
    budget = str(ws[BUDGET_CELL].value).strip()
    end = next(r for r in range(FIRST_DATA_ROW, ws.max_row + 1)
               if ws.cell(r, 2).value == SECTION_B_MARKER)
    day_cols = [c.column for c in ws[HEADER_ROW] if c.value == "# Days"]
    total_col, body = day_cols[-1], day_cols[:-1]
    coded = {c: str(ws.cell(CODE_ROW, c).value).strip() for c in body if ws.cell(CODE_ROW, c).value}

    for m in ws.merged_cells.ranges:
        if FIRST_DATA_ROW <= m.min_row < end:
            warns.append((budget, f"merged cells in staff section: {m}"))

    rows, recon_total, recon_excl = [], 0.0, 0.0
    for r in range(FIRST_DATA_ROW, end):
        role = ws.cell(r, ROLE_COL).value
        name = str(ws.cell(r, NAME_COL).value or "").strip()
        if name.upper().startswith(SUBTOTAL_PREFIX):
            continue
        if not name:
            for key, person in UNNAMED_ROW_OVERRIDES.items():
                if role and key in str(role):
                    name = person
        src_total = num(ws.cell(r, total_col).value)
        recon_total += src_total

        coded_sum = 0.0
        for c in body:
            v = ws.cell(r, c).value
            fv = wf.cell(r, c).value
            if isinstance(fv, str) and fv.startswith("=") and v is None:
                warns.append((budget, f"row {r}: formula in {col_letter(c)} has no cached value (open/save in Excel)"))
            if not num(v):
                continue
            if c not in coded:
                warns.append((budget, f"row {r}: {v} days in {col_letter(c)} has no milestone code"))
                continue
            coded_sum += v
            if name and not name.upper().startswith(EXCLUDE_NAME_PREFIXES):
                rows.append((name, budget, coded[c], round(v, 4), role))
        if abs(src_total - coded_sum) > 0.01:
            warns.append((budget, f"row {r} ({name or role}): milestone days {coded_sum:.2f} != source total {src_total:.2f}"))
        if not name or name.upper().startswith(EXCLUDE_NAME_PREFIXES):
            recon_excl += src_total
            if not name and coded_sum:  # TBD exclusions are intentional; unnamed rows need a human look
                warns.append((budget, f"row {r}: {coded_sum:.1f} days excluded - no name (position: {role})"))
    return rows, (budget, recon_total, recon_excl, fname)


def main(indir, outpath):
    warns, rows, recon = [], [], []
    for path in sorted(glob.glob(os.path.join(indir, "*.xlsx"))):
        if os.path.basename(path).startswith(SKIP_FILES) or os.path.basename(path).startswith("~$"):
            continue
        res = extract(path, warns)
        if res:
            rows += res[0]
            recon.append(res[1])

    totals = defaultdict(float)
    nbud = defaultdict(set)
    for p, b, _, d, _ in rows:
        totals[p] += d
        nbud[p].add(b)
    names = list(totals)
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            if difflib.SequenceMatcher(None, a.lower(), b.lower()).ratio() > 0.8:
                warns.append(("names", f"similar names, check for split totals: '{a}' / '{b}'"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LOE Table"
    ws.append(["Person", "Budget", "Milestone", "LOE (days)", "Role / Position", "Person Total LOE (all budgets)"])
    n = len(rows) + 1
    for i, r in enumerate(rows, 2):
        ws.append(list(r) + [f"=SUMIF($A$2:$A${n},A{i},$D$2:$D${n})"])
    t = Table(displayName="LOE", ref=f"A1:F{n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for col, w in zip("ABCDEF", [28, 20, 12, 12, 44, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    s = wb.create_sheet("Totals by Person")
    s.append(["Person", "Total LOE (days)", "# Budgets"])
    for p in sorted(totals, key=lambda p: -totals[p]):
        s.append([p, round(totals[p], 2), len(nbud[p])])
    for col, w in zip("ABC", [28, 16, 12]):
        s.column_dimensions[col].width = w
    s.freeze_panes = "A2"

    c = wb.create_sheet("Reconciliation")
    c.append(["Budget", "Source total-column days (excl. subtotal row)", "Excluded (TBD / unnamed)",
              "Expected in table", "Days in table", "Difference"])
    for b, tot, excl, _ in recon:
        i = c.max_row + 1
        c.append([b, round(tot, 3), round(excl, 3), f"=B{i}-C{i}",
                  f"=SUMIF('LOE Table'!$B:$B,A{i},'LOE Table'!$D:$D)", f"=ROUND(E{i}-D{i},3)"])
    i = c.max_row + 1
    c.append(["TOTAL"] + [f"=SUM({k}2:{k}{i - 1})" for k in "BCDEF"])
    for col, w in zip("ABCDEF", [22, 24, 22, 18, 14, 12]):
        c.column_dimensions[col].width = w

    w = wb.create_sheet("Warnings")
    w.append(["Source", "Warning"])
    for x in warns:
        w.append(list(x))
    w.column_dimensions["A"].width = 24
    w.column_dimensions["B"].width = 100

    wb.save(outpath)
    print(f"{len(rows)} rows, {len(totals)} people, {sum(totals.values()):.1f} days -> {outpath}")
    print(f"{len(warns)} warning(s)")
    for x in warns:
        print("  -", *x)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
