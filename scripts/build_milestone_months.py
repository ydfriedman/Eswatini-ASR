#!/usr/bin/env python3
"""Build the country | milestone | valid-months table.

Excel workplans (12 countries) are read programmatically from the Gantt tab
('Implementation Timeline').  The five Word core workplans (and the DRC
foundation milestones inside the DRC/Ebola doc) were read and cross-checked by
hand; their reviewed result is the CORE_ROWS table below, with the evidence
and conflicts recorded per row.

Usage: python3 scripts/build_milestone_months.py WORKPLAN_DIR OUTPUT.xlsx
"""
import glob, os, re, sys, warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
warnings.filterwarnings("ignore")

MON = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
def idx(y, m): return y * 12 + m - 1
def lab(i): return f"{MON[i % 12]}-{str(i // 12)[2:]}"
def span(a, b): return list(range(idx(*a), idx(*b) + 1))
S26, O26, D26, J27, M27, JN27, AU27, SE27 = (2026,9),(2026,10),(2026,12),(2027,1),(2027,3),(2027,6),(2027,8),(2027,9)
MONTHS_FROM_LABEL = {m: i + 1 for i, m in enumerate(MON)}

# ----------------------------------------------------------------- Excel part
def excel_rows(wpdir, overrides):
    rows, flags = [], []
    for f in sorted(glob.glob(os.path.join(wpdir, "*.xlsx"))):
        base = os.path.basename(f)
        country = re.match(r"AFGHSP_(.+?)(?:_Workplan|_ ?DOS)", base)[1].replace("_", " ")
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb["Implementation Timeline"]
        tab_no = wb.sheetnames.index("Implementation Timeline") + 1
        hdr = next(r for r in range(1, 12) if any(str(c.value or "").replace("\n", " ").strip().startswith("Month 1") for c in ws[r]))
        mcols = [c.column for c in ws[hdr] if str(c.value or "").startswith("Month")]
        idc = 2 if str(ws.cell(hdr, 1).value) == "Obj" else 1
        ttl = idc + 1
        first = str(ws.cell(hdr, mcols[0]).value).replace("\n", " ")
        m = re.search(r"Month 1 (\w{3})-(\d\d)", first)
        notes = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value and "MONTH ANCHORING" in str(c.value).upper())
        cover = " ".join(str(c.value) for row in wb["Cover Page"].iter_rows() for c in row if c.value and "Month 1 =" in str(c.value))
        if m: anchor, basis = (2000 + int(m[2]), MONTHS_FROM_LABEL[m[1]]), "Gantt column headers carry calendar months"
        elif "Month 1 = October 2026" in notes + cover: anchor, basis = O26, "Sheet note: Month 1 = October 2026"
        else: anchor, basis = S26, "ASSUMED: no calendar anchor in file; Sep-26 assumed (same template family / Sept 21 2026 start)"
        for r in range(hdr + 1, ws.max_row + 1):
            fill = ws.cell(r, idc).fill.fgColor.rgb
            a, t = ws.cell(r, idc).value, str(ws.cell(r, ttl).value or "")
            marks = [i + 1 for i, c in enumerate(mcols) if ws.cell(r, c).value not in (None, "")]
            if fill != "FFD6E4F0" or not marks: continue
            mm = re.match(r"^([A-Z]{3,5}\d{1,2})\s*[—-]\s*(.*)", t) or re.match(r"^(M[A-Z]{3}0)\s*[—-]?\s*(.*)", t)
            if mm: code, title = mm[1], mm[2]
            elif re.match(r"^[A-Z]{3,5}\d{1,2}$", str(a)): code, title = str(a), t
            else: continue
            months = [idx(*anchor) + k - 1 for k in marks]
            note = ""
            if (country, code) in overrides:
                months, note = overrides[(country, code)](anchor)
                flags.append((country, code, "Gantt vs. Planned End / Workplan tab", note))
            if anchor == S26 and basis.startswith("ASSUMED"):
                note = (note + " | " if note else "") + basis
            rows.append(dict(country=country, code=code, title=title.strip(), months=months,
                             typ="Mobilization (not in budget)" if re.match(r"^M[A-Z]{3}0$", code) else "Milestone",
                             source=f"{base} > tab {tab_no} 'Implementation Timeline' (Gantt)",
                             conf="High" if not note else "Medium", note=note))
    return rows, flags

def nam08(anchor):
    return span(O26, M27), ("Gantt bar ends Feb-27, but Planned End on the timeline tab and the Workplan tab both say Mar-27; "
                            "recommended window follows the two text sources.")
OVERRIDES = {("Namibia", "NAM08"): nam08}

# ------------------------------------------------------------------ Word part
# country, code, title, type, months, source, confidence, notes
G, H, SC, E, EB = "CORE G2G", "CORE-HIC", "CORE-Supply Chain", "CORE-EMR ECHIS", "CORE-Ebola"
CORE_ROWS = [
 (G,"MM0.3","Mobilization & operational readiness (PFM & G2G)","Mobilization (not in budget)",span(S26,D26),"G2G doc: milestone header + activity table (Sep-Dec 2026)","High",""),
 (G,"T1.5","G2G milestone frameworks, training & learning","Milestone",span(S26,SE27),"G2G doc: activity table + Gantt (T1.5-A1..A15)","Medium",
  "Header 'Start/End' and Planned Milestone Date say Aug-2027, and the Gantt title says 'FY2027 (Sep 2026 to Aug 2027)', but 5 activities (A1, A12-A15) and the Gantt run to Sep-2027. Followed activity-level evidence; if Aug-27 is right, drop Sep-27."),
 (SC,"MM0.4","Mobilization & operational readiness (Supply chain)","Mobilization (not in budget)",span(S26,D26),"Supply Chain doc: milestone header + activity table (Sep-Dec 2026)","High","No bar for MM0.4 in the doc's Gantt."),
 (SC,"T1.6","Digitally enabled diagnostics-to-treatment linkage","Milestone",span(O26,SE27),"Supply Chain doc: header (Month 1-12) + activity table + Gantt agree","High",""),
 (SC,"C5","Supply Chain Control Tower","Milestone",span(O26,SE27),"Supply Chain doc: activity table + Gantt agree (A1-A12)","High","Header describes build as Month 1-9 (to Jun-27) with sustainment/handover to Month 12; window shown includes sustainment."),
 (SC,"C6","eLMIS & last-mile supply chain digitization","Milestone",span(O26,SE27),"Supply Chain doc: header + activity table + Gantt agree","High",""),
 (SC,"C7","Molecular Diagnostics Control Tower","Milestone",span(O26,SE27),"Supply Chain doc: header + activity table + Gantt agree","High",""),
 (H,"MM0.5","Mobilization & operational readiness (HIC)","Mobilization (not in budget)",span(S26,D26),"HIC doc: milestone header + activity table (Sep-Dec 2026)","High","No bar for MM0.5 in the doc's Gantt."),
 (H,"C1a","Generalized HIC Digital Kit (global good)","Milestone",span(S26,M27),"HIC doc: activity table (Sep-26..Mar-27); header says Month 1 (Oct-26)..Month 6 (Mar-27); Planned Milestone Date 31 Mar 2027","Medium",
  "GANTT DISAGREES: the C1a bars are compressed into Sep-Dec 2026 (e.g. A12 publication shown Dec-26 vs Mar-27 in the table; A4-A11 start/end 1-2 months early), plus a Jan-Aug 2027 'kit maintenance' row (A13) that has no activity-table row (the text assigns maintenance to C1b). Followed the activity table. Header start is Oct-26; A1-A3 begin Sep-26."),
 (H,"C1b","HIC country implementations (Kenya, Botswana, Sierra Leone)","Milestone",span(S26,AU27),"HIC doc: activity table + Gantt agree (Sep-26..Aug-27)","Medium",
  "Header says Month 1 (Oct-26)..Month 11 (Aug-27) while the progress-review table lists 'Final Achievement' at Month 12 (Sep-27); A1 starts Sep-26. Gantt row A11 has a Jul-27 gap (Jun and Aug only) vs. table Jun-Aug. Table lists activity 'C1b-A17' (likely A7)."),
 (E,"MM0.1","Mobilization & operational readiness (C2)","Mobilization (not in budget)",span(S26,D26),"EMR/C2/T1.8 doc: header + activity table (Sep-Dec 2026)","High",""),
 (E,"MM0.2","Mobilization & operational readiness (T1.8)","Mobilization (not in budget)",span(S26,D26),"EMR/C2/T1.8 doc: header + activity table (Sep-Dec 2026)","High",""),
 (E,"C2","eCHIS global public good","Milestone",span(O26,SE27),"EMR/C2/T1.8 doc: 'Period Month 1-Month 12' + activity table + Gantt agree (M1-M12)","Medium",
  "Doc never states what calendar month 'Month 1' is. Month 1 = Oct-26 is inferred: Q3 reviews at M7-M9 (Apr-Jun) and the T1.8/C2 planned milestone date of 14 Sep 2027 only fit Oct-26 start."),
 (E,"T1.8","High-impact HIV interventions (Kenya, Eswatini)","Milestone",span(O26,SE27),"EMR/C2/T1.8 doc: 'Period Month 1-Month 12' + activity table + Gantt agree (M1-M12)","Medium",
  "Same Month 1 = Oct-26 inference as C2. Minor: T1.8-A05 table says M1-M2, Gantt shows M2 only."),
 (EB,"MM0.6","Mobilization & operational readiness (Ebola/DRC)","Mobilization (not in budget)",span(O26,D26),"DRC/EVD doc: header (Oct-Dec 2026) + activity table + both Gantts agree","High",""),
 (EB,"T1.4","Faster EVD detection, notification & response (7-1-7)","Milestone",span(J27,SE27),"DRC/EVD doc: header + activity table + both Gantts agree","High",""),
 (EB,"E1","EVD preparedness, early warning, surge staffing & contact tracing","Milestone",span(J27,JN27),"DRC/EVD doc: activity table + detailed Gantt agree (Jan-Jun 2027); evidence due Q3 FY2027","Medium",
  "CONFLICT: milestone header End Month and Planned Milestone Date say Sep-2027 (Q4 FY2027), and the summary Gantt shades Jan-Sep, but every activity ends by Jun-27 and all E1 evidence is due Q3 FY2027. Window shown = activity window; extend to Sep-27 if the header governs."),
 (EB,"E2","Cross-border EVD detection & response coordination","Milestone",span(J27,SE27),"DRC/EVD doc: activity table + detailed Gantt + summary Gantt agree (Jan-Sep 2027)","Medium",
  "CONFLICT: milestone header Start Month says October 2026, but the first activity (E2-A1.1) starts Jan-27 in the table and both Gantts. Followed the activities."),
 (EB,"E3","EVD laboratory connectivity & case-management data readiness","Milestone",span(O26,SE27),"DRC/EVD doc: header (Month 1-12) + activity table + detailed Gantt agree","Medium",
  "SUMMARY GANTT DISAGREES: it shades E3 Jan-Sep 2027 like the other EVD milestones; the header, the activity table (E3-A1 starts M1) and the detailed Gantt all start Oct-26."),
]
# DRC 'foundation' milestones: only in the DRC/EVD doc's Gantts (summary and detailed agree), no activity table
DRC_FOUND = [("DRC00","Mobilization","Mobilization (not in budget)",span(O26,D26),"Q1 FY2027"),
 ("DRC01","Cybersecurity","Milestone",span(O26,(2027,2)),"Q2 FY2027"),
 ("DRC02","DHIS2 Polio/AFP","Milestone",span(O26,M27),"Q3 FY2027"),
 ("DRC03","Data-quality supervision","Milestone",span(O26,M27),"Q3 FY2027"),
 ("DRC04","Interoperability / eLIS","Milestone",span(O26,SE27),"Q4 FY2027"),
 ("DRC05","EMR site readiness","Milestone",span(O26,SE27),"Q4 FY2027"),
 ("DRC06","OpenMRS deployment","Milestone",span(O26,SE27),"Q4 FY2027"),
 ("DRC07","Data warehouse","Milestone",span(O26,SE27),"Q4 FY2027"),
 ("DRC08","DQR cycles","Milestone",span(O26,SE27),"Q4 FY2027"),
 ("DRC09","Digital governance","Milestone",span(O26,SE27),"Q4 FY2027")]

def main(wpdir, out):
    rows, flags = excel_rows(wpdir, OVERRIDES)
    for c, code, title, typ, months, src, conf, note in CORE_ROWS:
        rows.append(dict(country=c, code=code, title=title, months=months, typ=typ, source=src, conf=conf, note=note))
        if note and conf != "High": flags.append((c, code, "Word doc cross-check", note))
    for code, title, typ, months, q in DRC_FOUND:
        rows.append(dict(country="DRC", code=code, title=title, months=months, typ=typ,
                         source="DRC/EVD doc: 'DRC AFGHS-P Foundation' rows of the summary and detailed Gantts (they agree)",
                         conf="Medium", note=f"Gantt-only: the doc has no activity table or header dates for DRC00-DRC09 to cross-check. Payment-quarter target {q}."))
    rows.sort(key=lambda r: (r["country"], r["code"]))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Milestone Months"
    ws.append(["Country / Budget", "Milestone", "Valid months", "First month", "Last month", "# months", "Type", "Title", "Confidence", "Notes / conflicts", "Source"])
    for r in rows:
        m = r["months"]; contiguous = m == list(range(m[0], m[-1] + 1))
        txt = f"{lab(m[0])} to {lab(m[-1])}" if contiguous else ", ".join(lab(x) for x in m)
        ws.append([r["country"], r["code"], txt, lab(m[0]), lab(m[-1]), len(m), r["typ"], r["title"], r["conf"], r["note"], r["source"]])
    t = Table(displayName="MilestoneMonths", ref=f"A1:K{ws.max_row}"); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); ws.add_table(t)
    for col, w in zip("ABCDEFGHIJK", [20, 11, 20, 11, 11, 9, 26, 52, 11, 90, 60]): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[9].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    lg = wb.create_sheet("Monthly (long)"); lg.append(["Country / Budget", "Milestone", "Month", "Month (date)", "Type", "Confidence"])
    import datetime
    for r in rows:
        for x in r["months"]: lg.append([r["country"], r["code"], lab(x), datetime.date(x // 12, x % 12 + 1, 1), r["typ"], r["conf"]])
    for row in lg.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "mmm-yy"
    t2 = Table(displayName="MonthlyLong", ref=f"A1:F{lg.max_row}"); t2.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); lg.add_table(t2)
    for col, w in zip("ABCDEF", [20, 11, 10, 13, 26, 11]): lg.column_dimensions[col].width = w
    lg.freeze_panes = "A2"
    fl = wb.create_sheet("Discrepancies"); fl.append(["Country / Budget", "Milestone", "Check", "What disagrees / what I did"])
    for x in flags: fl.append(list(x))
    for col, w in zip("ABCD", [20, 11, 34, 140]): fl.column_dimensions[col].width = w
    for row in fl.iter_rows(min_row=2): row[3].alignment = Alignment(wrap_text=True, vertical="top")
    cv = wb.create_sheet("Coverage"); cv.append(["Budget", "Status", "Detail"])
    cv.append(["Honduras", "NO WORKPLAN PROVIDED", "Budget has HON01-HON04; no workplan file in the upload."])
    cv.append(["Sierra Leone", "NO WORKPLAN PROVIDED", "Budget has SIL01-SIL03; no workplan file in the upload. (The HIC doc lists Sierra Leone only as a C1b implementation country.)"])
    cv.append(["DRC", "Partial source", "No DRC country workplan file; DRC00-DRC09 months come only from the 'DRC AFGHS-P Foundation' rows of the DRC/EVD Word doc's Gantts."])
    cv.append(["All other budgets", "Covered", "Every budget milestone column has a valid-months row (mobilization rows are extra, with no budget column)."])
    for col, w in zip("ABC", [20, 26, 130]): cv.column_dimensions[col].width = w
    wb.save(out)
    print(f"{len(rows)} milestone rows -> {out}; {len(flags)} flagged")

if __name__ == "__main__":
    if len(sys.argv) != 3: sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
